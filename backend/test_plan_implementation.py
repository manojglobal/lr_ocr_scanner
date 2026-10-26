"""
Test Script: Verification of Himalaya LR Plan Implementation
Tests:
1. 'BNF LR No' removal from schema and mapping.
2. Primary 'LR No *' / 'grno (*)' extraction and mapping.
3. Container Seal Number mapping to 'pmark' (and never into LR fields).
4. 'driver_master.json' auto-fetch is disabled: absent driver details remain strictly empty.
5. Driver details present in document packet are properly retained.
6. Excel workbook generation produces clean 49-column output matching BNF file Modify.xlsx.
7. Verification of prompt constraints in DEFAULT_HANDWRITTEN_OCR_PROMPT.
"""
import sys
import os
import unittest
import openpyxl
import io

# Setup path to backend
BACKEND_DIR = os.path.dirname(os.path.abspath(__file__))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)
os.chdir(BACKEND_DIR)

import main


class TestHimalayaPlanImplementation(unittest.TestCase):

    def test_01_headers_count_and_no_bnf_lr(self):
        """Verify headers match 49 columns and do not contain BNF LR No."""
        headers = main.GR_BNF_49_HEADERS
        print(f"\n[Test 1] Header count: {len(headers)}")
        self.assertEqual(len(headers), 49, f"Expected 49 columns, found {len(headers)}")
        self.assertNotIn("BNF LR No", headers, "'BNF LR No' should NOT be present in 49-column headers")
        self.assertIn("grno (*)", headers, "'grno (*)' must be present")
        self.assertIn("pmark", headers, "'pmark' must be present for container seal number")
        self.assertIn("Vehicle No *", headers, "'Vehicle No *' must be present")
        self.assertIn("Driver Name *", headers, "'Driver Name *' must be present")

    def test_02_map_to_gr_row_primary_lr_and_seal_number(self):
        """Verify primary LR maps to grno (*), seal number maps to pmark, and BNF LR No is omitted."""
        test_data = {
            "LR No *": "32978",
            "grno (*)": "32978",
            "grdt *": "30/09/2026",
            "Customer Name *": "Himalaya Wellness Company",
            "Vehicle No *": "GJ05BZ0237",
            "Vehicle Type *": "32FT CONTAINER",
            "orgrefcode": "Bangalore (MWH)",
            "destrefcode": "Indore (MWH)",
            "cngr *": "M/s Himalaya Wellness Company",
            "cnge *": "M/s Himalaya Wellness Company Indore",
            "pmark": "00153",  # Container seal number
            "bookedpckgs": "2812",
            "packing": "Box",
            "goods": "AYURVEDIC MEDICINES",
            "valgoods": "11488522.00",
            "invoiceno": "5526701832",
            "ewaybillno": "102566549518",
            # Intentionally no driver information:
            "Driver Name *": "",
            "Mobile no   *": "",
            "Liscense No *": "",
            "Valid Up to *": "",
        }

        row = main.map_to_gr_row(test_data, main.GR_BNF_49_HEADERS)
        self.assertEqual(len(row), 49, f"Row length must be 49, got {len(row)}")

        # Check grno (*) index (col 1)
        grno_idx = main.GR_BNF_49_HEADERS.index("grno (*)")
        self.assertEqual(str(row[grno_idx]), "32978")

        # Check pmark index (col 16) - container seal number preserves leading zeros
        pmark_idx = main.GR_BNF_49_HEADERS.index("pmark")
        self.assertEqual(str(row[pmark_idx]), "00153")

        # Check that driver fields remain empty
        driver_idx = main.GR_BNF_49_HEADERS.index("Driver Name *")
        self.assertIn(row[driver_idx], (None, "", "None"))
        mobile_idx = main.GR_BNF_49_HEADERS.index("Mobile no   *")
        self.assertIn(row[mobile_idx], (None, "", "None"))

    def test_03_driver_master_auto_fetch_is_disabled(self):
        """
        Verify that a vehicle known in driver_master.json (e.g. KA23B5235) does NOT
        auto-populate driver details when those details are absent from the document.
        """
        # KA23B5235 has DHANAJI S PATIL in driver_master.json
        test_data = {
            "LR No *": "33118",
            "Vehicle No *": "KA23B5235",
            # Driver details are absent in document:
            "Driver Name *": "",
            "Mobile no   *": "",
            "Liscense No *": "",
            "Valid Up to *": "",
        }

        row = main.map_to_gr_row(test_data, main.GR_BNF_49_HEADERS)
        driver_idx = main.GR_BNF_49_HEADERS.index("Driver Name *")
        mobile_idx = main.GR_BNF_49_HEADERS.index("Mobile no   *")
        license_idx = main.GR_BNF_49_HEADERS.index("Liscense No *")

        # Must NOT be DHANAJI S PATIL
        self.assertNotEqual(row[driver_idx], "DHANAJI S PATIL", "Driver Name should NOT be auto-filled from registry")
        self.assertIn(row[driver_idx], (None, "", "None"), "Driver Name must remain empty for manual entry")
        self.assertIn(row[mobile_idx], (None, "", "None"), "Mobile number must remain empty")
        self.assertIn(row[license_idx], (None, "", "None"), "License number must remain empty")

    def test_04_driver_details_preserved_when_present_in_document(self):
        """Verify that driver details extracted from an attached DL or document ARE preserved."""
        test_data = {
            "LR No *": "32978",
            "Vehicle No *": "GJ05BZ0237",
            "Driver Name *": "JAYMANGAL RAI",
            "Mobile no   *": "9691474127",
            "Liscense No *": "MH02 20040039639",
            "Valid Up to *": "27/08/2030",
        }

        row = main.map_to_gr_row(test_data, main.GR_BNF_49_HEADERS)
        driver_idx = main.GR_BNF_49_HEADERS.index("Driver Name *")
        mobile_idx = main.GR_BNF_49_HEADERS.index("Mobile no   *")
        license_idx = main.GR_BNF_49_HEADERS.index("Liscense No *")
        valid_idx = main.GR_BNF_49_HEADERS.index("Valid Up to *")

        self.assertEqual(row[driver_idx], "JAYMANGAL RAI")
        self.assertEqual(str(row[mobile_idx]), "9691474127")
        self.assertEqual(row[license_idx], "MH02 20040039639")
        self.assertEqual(row[valid_idx], "27/08/2030")

    def test_05_prompt_rules_verification(self):
        """Verify the OCR prompt contains the correct zero-hallucination and empty field rules."""
        prompt = main.DEFAULT_HANDWRITTEN_OCR_PROMPT
        self.assertNotIn("treat it as BNF LR No", prompt)
        self.assertIn("STRICT EMPTY FIELD RULE", prompt)
        self.assertIn("pmark", prompt)
        self.assertIn("Extract ONLY the actual seal code or digits", prompt)
        self.assertIn("LR No *", prompt)
        self.assertIn("CONSIGNMENT NOTE: No.", prompt)

    def test_06_workbook_generation_structure(self):
        """Verify workbook generation builds 49 columns without BNF LR No."""
        doc1 = main.ParsedDocument(
            document_title="LR_32978",
            data={
                "LR No *": "32978",
                "grno (*)": "32978",
                "grdt *": "30/09/2026",
                "Customer Name *": "Himalaya Wellness Company",
                "Vehicle No *": "GJ05BZ0237",
                "Driver Name *": "JAYMANGAL RAI",
                "pmark": "00153",
            },
            key_value_pairs=[],
            confidence_notes=""
        )

        doc2 = main.ParsedDocument(
            document_title="LR_33118",
            data={
                "LR No *": "33118",
                "grno (*)": "33118",
                "grdt *": "02/09/2026",
                "Vehicle No *": "MP09HH0830",
                # Empty driver info:
                "Driver Name *": "",
            },
            key_value_pairs=[],
            confidence_notes=""
        )

        user_wb, user_bytes, _, _ = main._build_workbook_for_docs(
            [doc1, doc2],
            format_type="gr_bnf_49",
            user_name="QA Tester",
            include_audit_meta=False
        )

        wb = openpyxl.load_workbook(io.BytesIO(user_bytes))
        ws = wb.active

        # Check row 1 headers
        header_vals = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(len(header_vals), 49)
        self.assertNotIn("BNF LR No", header_vals)
        self.assertIn("grno (*)", header_vals)
        self.assertIn("pmark", header_vals)

        # Check row 2 (doc1)
        grno_col = header_vals.index("grno (*)") + 1
        pmark_col = header_vals.index("pmark") + 1
        driver_col = header_vals.index("Driver Name *") + 1

        self.assertEqual(str(ws.cell(row=2, column=grno_col).value), "32978")
        self.assertEqual(str(ws.cell(row=2, column=pmark_col).value), "00153")
        self.assertEqual(str(ws.cell(row=2, column=driver_col).value), "JAYMANGAL RAI")

        # Check row 3 (doc2 - absent driver must be empty)
        self.assertEqual(str(ws.cell(row=3, column=grno_col).value), "33118")
        self.assertIn(ws.cell(row=3, column=driver_col).value, (None, "", "None"))

    def test_07_transit_days_calculation(self):
        """
        Verify 'transit_days' calculation:
        - integer difference in days between 'grdt' and 'Valid Upto'
        - if either date is missing, unparseable, or invalid -> empty string / None (no fallback 2)
        """
        # Case A: 30/09/2026 to 08/10/2026 -> 8 days
        data_a = {
            "grdt *": "30/09/2026",
            "Valid Upto": "08/10/2026",
        }
        row_a = main.map_to_gr_row(data_a, main.GR_BNF_49_HEADERS)
        transit_idx = main.GR_BNF_49_HEADERS.index("Transit Days")
        self.assertEqual(row_a[transit_idx], 8)

        # Case B: 08/10/2026 to 18/10/2026 -> 10 days
        data_b = {
            "grdt *": "08/10/2026",
            "Valid Upto": "18/10/2026",
        }
        row_b = main.map_to_gr_row(data_b, main.GR_BNF_49_HEADERS)
        self.assertEqual(row_b[transit_idx], 10)

        # Case C: Missing Valid Upto -> must be empty (NOT default 2)
        data_c = {
            "grdt *": "30/09/2026",
            "Valid Upto": "",
        }
        row_c = main.map_to_gr_row(data_c, main.GR_BNF_49_HEADERS)
        self.assertIn(row_c[transit_idx], (None, "", "None"))
        self.assertNotEqual(row_c[transit_idx], 2, "Should NOT default to 2 when Valid Upto is missing")

        # Case D: Missing grdt -> must be empty (NOT default 2)
        data_d = {
            "grdt *": "",
            "Valid Upto": "08/10/2026",
        }
        row_d = main.map_to_gr_row(data_d, main.GR_BNF_49_HEADERS)
        self.assertIn(row_d[transit_idx], (None, "", "None"))
        self.assertNotEqual(row_d[transit_idx], 2, "Should NOT default to 2 when grdt is missing")

        # Case E: Invalid/Unparseable date -> must be empty (NOT default 2)
        data_e = {
            "grdt *": "invalid_date",
            "Valid Upto": "not_a_date",
        }
        row_e = main.map_to_gr_row(data_e, main.GR_BNF_49_HEADERS)
        self.assertIn(row_e[transit_idx], (None, "", "None"))
        self.assertNotEqual(row_e[transit_idx], 2, "Should NOT default to 2 when date is invalid")

    def test_08_pmark_cleaning_and_igst_zero(self):
        """
        Verify:
        - 'pmark' Cleaning Rule: Extract ONLY the actual seal code or digits (e.g. CJ007567 or 00153).
          Strictly strip and ignore any prefixes, labels, or symbols like Seal, Sealv0, No, =, or :
        - 'igst': Always output 0.00 or 0 for manual data entry purposes.
        """
        # Test cleaning helper directly
        self.assertEqual(main._clean_pmark("Sealv0 = CJ007567"), "CJ007567")
        self.assertEqual(main._clean_pmark("Seal No 00153"), "00153")
        self.assertEqual(main._clean_pmark("Seal: 00105"), "00105")
        self.assertEqual(main._clean_pmark("Seal# 99887"), "99887")
        self.assertEqual(main._clean_pmark("= 12345 ="), "12345")

        # Test map_to_gr_row with various pmark formats and verify igst = 0.00
        test_cases = [
            ("Sealv0 = CJ007567", "CJ007567"),
            ("Seal No 00153", "00153"),
            ("Seal: 00105", "00105"),
        ]
        pmark_idx = main.GR_BNF_49_HEADERS.index("pmark")
        igst_idx = main.GR_BNF_49_HEADERS.index("igst")

        for raw_pmark, expected in test_cases:
            data = {
                "LR No *": "32978",
                "pmark": raw_pmark,
                "igst": "1800.50",  # Even if raw data had a tax value, igst must output 0.00
            }
            row = main.map_to_gr_row(data, main.GR_BNF_49_HEADERS)
            self.assertEqual(str(row[pmark_idx]), expected)
            self.assertEqual(str(row[igst_idx]), "0.00")

    def test_09_grno_capture_without_dropping(self):
        """
        Verify:
        - grno / LR No *: Extract the Consignment Note number from the 'CONSIGNMENT NOTE: No.' box.
          If it's a 5-digit number or any valid consignment serial, capture it accurately without dropping it.
        """
        # Standalone 5-digit numbers
        self.assertEqual(main._clean_grno("32760"), "32760")
        self.assertEqual(main._clean_grno("30135"), "30135")
        self.assertEqual(main._clean_grno("32754"), "32754")
        self.assertEqual(main._clean_grno("CONSIGNMENT NOTE: No. 32760"), "32760")
        self.assertEqual(main._clean_grno("LR 33118"), "33118")
        self.assertEqual(main._clean_grno("TG30135"), "30135")
        self.assertEqual(main._clean_grno("TG-32754"), "32754")

        # Non-5-digit valid consignment serials must NOT be dropped or truncated
        self.assertEqual(main._clean_grno("9421736"), "9421736")
        self.assertEqual(main._clean_grno("CN-987654"), "CN-987654")
        self.assertEqual(main._clean_grno("TG10849828"), "TG10849828")

        # Verify mapping when only 'LR No *' is provided in input data
        grno_idx = main.GR_BNF_49_HEADERS.index("grno (*)")
        data_lr_only = {
            "LR No *": "32760",
        }
        row_lr = main.map_to_gr_row(data_lr_only, main.GR_BNF_49_HEADERS)
        self.assertEqual(str(row_lr[grno_idx]), "32760")

        # Verify mapping when valid serial is provided
        data_serial = {
            "LR No *": "TG10849828",
        }
        row_serial = main.map_to_gr_row(data_serial, main.GR_BNF_49_HEADERS)
        self.assertEqual(str(row_serial[grno_idx]), "TG10849828")

    def test_10_patanjali_prompt_rules(self):
        """Verify the Patanjali OCR prompt contains the correct margin LR rules and FTL charge rules."""
        prompt = main.DEFAULT_PATANJALI_OCR_PROMPT
        self.assertIn("PATANJALI PARIVAHAN", prompt)
        self.assertIn("top-left and top-right", prompt)
        self.assertIn("pencil or pen", prompt)
        self.assertIn("leave it strictly empty", prompt)
        self.assertIn("DO NOT output \"0\"", prompt)
        self.assertIn("Leaving the field empty ensures the record cannot be saved", prompt)
        self.assertIn("DO NOT fall back to the printed 'C Note No.'", prompt)
        self.assertIn("NEVER copy or confuse vehicle registration numbers", prompt)
        self.assertIn("xxxx", prompt)

    def test_11_patanjali_margin_lr_capture(self):
        """Verify that Patanjali independent margin handwritten LR (e.g. 21349) maps correctly into grno (*)."""
        data = {
            "Customer Name *": "PATANJALI FOODS LTD BISCUIT (Formerly RSIL)",
            "grno (*)": "21349",
            "LR No *": "21349",
            "grdt *": "05/10/2026",
            "Vehicle No *": "RJ32GC5386",
            "Vehicle Type *": "Market",
            "freight  *": "0",
            "cgst": "0.00",
            "sgst": "0.00",
            "igst": "0.00",
        }
        row = main.map_to_gr_row(data, main.GR_BNF_49_HEADERS)
        grno_idx = main.GR_BNF_49_HEADERS.index("grno (*)")
        self.assertEqual(str(row[grno_idx]), "21349")

    def test_12_patanjali_clean_margin_leaves_empty_so_cannot_save(self):
        """
        Verify: If nothing is written with a pencil or pen in the top-left or top-right section of a Patanjali bilty,
        it leaves grno (*) and LR No * strictly EMPTY (None / "") - never displaying 0.
        This ensures mandatory field validation blocks saving until manually entered.
        """
        data = {
            "Customer Name *": "PATANJALI FOODS LTD BISCUIT (Formerly RSIL)",
            "grno (*)": "",
            "LR No *": "",
            "grdt *": "09/10/2026",
            "Vehicle No *": "BR03K4305",
            "Vehicle Type *": "Market",
        }
        row = main.map_to_gr_row(data, main.GR_BNF_49_HEADERS)
        grno_idx = main.GR_BNF_49_HEADERS.index("grno (*)")
        self.assertIn(row[grno_idx], (None, "", "None"), "Clean margin Patanjali bilty must be left empty in grno (*)")
        self.assertNotEqual(str(row[grno_idx]), "0", "Must NEVER display '0' for missing Patanjali LR")
        self.assertNotEqual(str(row[grno_idx]), "9421832", "Must NEVER fall back to printed C Note No")

        # Also test if data had "0", it gets cleaned to empty (None) for Patanjali
        data_zero = {
            "Customer Name *": "PATANJALI FOODS LTD BISCUIT (Formerly RSIL)",
            "grno (*)": "0",
            "LR No *": "0",
            "grdt *": "09/10/2026",
            "Vehicle No *": "BR03K4305",
        }
        row_zero = main.map_to_gr_row(data_zero, main.GR_BNF_49_HEADERS)
        self.assertIn(row_zero[grno_idx], (None, "", "None"), "Even if '0' was present, Patanjali must map to empty/None")
        self.assertNotEqual(str(row_zero[grno_idx]), "0", "Must NOT output '0'")

    def test_13_patanjali_ftl_charges_mapping(self):
        """Verify that Patanjali FTL charges map to 0/0.00."""
        data = {
            "Customer Name *": "PATANJALI FOODS LTD",
            "grno (*)": "21348",
            "LR No *": "21348",
            "freight  *": 0,
            "cgst": 0.0,
            "sgst": 0.0,
            "igst": "0.00",
        }
        row = main.map_to_gr_row(data, main.GR_BNF_49_HEADERS)
        fr_idx = main.GR_BNF_49_HEADERS.index("freight  *")
        cgst_idx = main.GR_BNF_49_HEADERS.index("cgst")
        sgst_idx = main.GR_BNF_49_HEADERS.index("sgst")
        igst_idx = main.GR_BNF_49_HEADERS.index("igst")
        self.assertEqual(row[fr_idx], 0)
        self.assertEqual(row[cgst_idx], 0)
        self.assertEqual(row[sgst_idx], 0)
        self.assertEqual(str(row[igst_idx]), "0.00")

    def test_14_workbook_generation_with_patanjali_and_himalaya(self):
        """Verify clean 49-column workbook generation with both Himalaya and Patanjali documents."""
        doc_himalaya = main.ParsedDocument(
            document_title="Himalaya_LR_32978",
            data={
                "LR No *": "32978",
                "grno (*)": "32978",
                "grdt *": "30/09/2026",
                "Customer Name *": "Himalaya Wellness Company",
                "Vehicle No *": "GJ05BZ0237",
                "pmark": "00153",
            },
            key_value_pairs=[],
            confidence_notes=""
        )
        doc_patanjali = main.ParsedDocument(
            document_title="Patanjali_LR_21349",
            data={
                "LR No *": "21349",
                "grno (*)": "21349",
                "grdt *": "05/10/2026",
                "Customer Name *": "PATANJALI FOODS LTD BISCUIT (Formerly RSIL)",
                "Vehicle No *": "RJ32GC5386",
                "pmark": "",
                "freight  *": "0",
            },
            key_value_pairs=[],
            confidence_notes=""
        )
        doc_patanjali_zero = main.ParsedDocument(
            document_title="Patanjali_LR_Clean_Margin",
            data={
                "LR No *": "0",
                "grno (*)": "0",
                "grdt *": "09/10/2026",
                "Customer Name *": "PATANJALI FOODS LTD",
                "Vehicle No *": "BR03K4305",
                "freight  *": "0",
            },
            key_value_pairs=[],
            confidence_notes=""
        )

        _, user_bytes, _, _ = main._build_workbook_for_docs(
            [doc_himalaya, doc_patanjali, doc_patanjali_zero],
            format_type="gr_bnf_49",
            user_name="Multi-Brand QA Tester",
            include_audit_meta=False
        )

        wb = openpyxl.load_workbook(io.BytesIO(user_bytes))
        ws = wb.active

        header_vals = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(len(header_vals), 49)

        grno_col = header_vals.index("grno (*)") + 1
        # Row 2 (Himalaya): 32978
        self.assertEqual(str(ws.cell(row=2, column=grno_col).value), "32978")
        # Row 3 (Patanjali with margin sequence): 21349
        self.assertEqual(str(ws.cell(row=3, column=grno_col).value), "21349")
        # Row 4 (Patanjali with clean margin): empty cell (None), never displaying 0
        self.assertIn(ws.cell(row=4, column=grno_col).value, (None, "", "None"))
        self.assertNotEqual(str(ws.cell(row=4, column=grno_col).value), "0")

    def test_15_grdt_formatting_strict_ddmmyyyy(self):
        """
        Verify:
        - Dates formatted as '12-OCT-2026', '12-Oct-26', '05 Oct 26', '2026-10-12'
          are strictly formatted as 'DD/MM/YYYY' (e.g. '12/10/2026', '05/10/2026').
        - Never show month names like 12-OCT-2026.
        """
        # Test helper directly
        self.assertEqual(main._format_date_ddmmyyyy("12-OCT-2026"), "12/10/2026")
        self.assertEqual(main._format_date_ddmmyyyy("12-Oct-26"), "12/10/2026")
        self.assertEqual(main._format_date_ddmmyyyy("05 Oct 26"), "05/10/2026")
        self.assertEqual(main._format_date_ddmmyyyy("2026-10-12"), "12/10/2026")
        self.assertEqual(main._format_date_ddmmyyyy("12/10/2026"), "12/10/2026")

        # Test map_to_gr_row with '12-OCT-2026' in grdt *
        data = {
            "LR No *": "32978",
            "grdt *": "12-OCT-2026",
            "Valid Upto": "18-OCT-2026",
            "Loading Start date *": "12-OCT-2026",
            "invoiedt": "12-OCT-2026",
        }
        row = main.map_to_gr_row(data, main.GR_BNF_49_HEADERS)
        grdt_idx = main.GR_BNF_49_HEADERS.index("grdt *")
        loading_idx = main.GR_BNF_49_HEADERS.index("Loading Start date *")
        valid_idx = main.GR_BNF_49_HEADERS.index("Valid Upto")
        invdt_idx = main.GR_BNF_49_HEADERS.index("invoiedt")
        transit_idx = main.GR_BNF_49_HEADERS.index("Transit Days")

        self.assertEqual(str(row[grdt_idx]), "12/10/2026")
        self.assertEqual(str(row[loading_idx]), "12/10/2026")
        self.assertEqual(str(row[valid_idx]), "18/10/2026")
        self.assertEqual(str(row[invdt_idx]), "12/10/2026")
        self.assertEqual(row[transit_idx], 6)

    def test_16_clean_pmark_preserves_leading_zeros_and_strips_middle_and_margin_seal_labels(self):
        """
        Verify that _clean_pmark strips prefixes and typos (Seal., Seal No-, Seal no:, Seaimot)
        while strictly preserving leading zeros in both margin and middle-of-text seal numbers.
        """
        # Margin seal formats
        self.assertEqual(main._clean_pmark("Seal. No- 03386"), "03386")
        self.assertEqual(main._clean_pmark("Seal No - 03330"), "03330")
        self.assertEqual(main._clean_pmark("Seal. No- 03301"), "03301")
        # Middle-of-text seal formats (embedded in goods or LR text)
        self.assertEqual(main._clean_pmark("Seal no: 002765"), "002765")
        self.assertEqual(main._clean_pmark("Seaimot 002765"), "002765")
        self.assertEqual(main._clean_pmark("Seuinoł 002765"), "002765")
        self.assertEqual(main._clean_pmark("Seuinol 002765"), "002765")
        # Alphanumeric seals and single-number seals
        self.assertEqual(main._clean_pmark("CJ007567"), "CJ007567")
        self.assertEqual(main._clean_pmark("Seal: CJ007567"), "CJ007567")
        self.assertEqual(main._clean_pmark("00153"), "00153")
        self.assertEqual(main._clean_pmark("Seal = 00153"), "00153")
        self.assertEqual(main._clean_pmark("0"), "0")
        self.assertEqual(main._clean_pmark(""), "")
        self.assertEqual(main._clean_pmark(None), "")

    def test_17_unicharm_prompt_content_and_rules(self):
        """Verify the Unicharm OCR prompt contains all crucial rules for Vuicharm/Unicharm."""
        prompt = main.DEFAULT_UNICHARM_OCR_PROMPT
        self.assertIn("UNICHARM INDIA PVT. LTD.", prompt)
        self.assertIn("BOX N FREIGHT LOGISTICS SOLUTIONS PVT. LTD.", prompt)
        self.assertIn("CONSIGNMENT NOTE: No.", prompt)
        self.assertIn("pmark", prompt)
        self.assertIn("PROMPT CASE - SEAL NUMBER IN THE MIDDLE OF THE \"LR\" IMAGE", prompt)
        self.assertIn("33306.pdf", prompt)
        self.assertIn("Description (Said to Contain)", prompt)
        self.assertIn("NEVER extract or confuse this seal number or description text as the LR number", prompt)
        self.assertIn("STRICT EMPTY FIELD RULE", prompt)
        self.assertIn("DD/MM/YYYY", prompt)

    def test_18_map_to_gr_row_unicharm_packet(self):
        """
        Verify mapping a Unicharm packet:
        - 5-digit LR number (33306)
        - Middle-of-text seal number (002765) cleanly in pmark
        - Absent driver details strictly empty
        - igst = 0.00
        - grdt strictly formatted DD/MM/YYYY
        """
        data = {
            "LR No *": "33306",
            "grno (*)": "33306",
            "grdt *": "29/09/2026",
            "Customer Name *": "UNICHARM INDIA PVT. LTD.",
            "Vehicle No *": "RJ18GB7132",
            "Vehicle Type *": "Market",
            "orgrefcode": "SANAND",
            "destrefcode": "AMBALA",
            "cngr *": "UNICHARM INDIA PVT. LTD., Sanand",
            "cnge *": "GLOBAL MARKETING, Ambala",
            "pmark": "Seaimot 002765",  # Middle-of-LR-text seal notation
            "bookedpckgs": "797",
            "packing": "Box",
            "goods": "Baby Diapers",
            "valgoods": "2112307.00",
            "invoiceno": "3324726-000484",
            "ewaybillno": "392344317558",
            "Driver Name *": "",
            "Mobile no   *": "",
            "Liscense No *": "",
            "Valid Up to *": "",
        }
        row = main.map_to_gr_row(data, main.GR_BNF_49_HEADERS)
        self.assertEqual(len(row), 49)

        grno_idx = main.GR_BNF_49_HEADERS.index("grno (*)")
        pmark_idx = main.GR_BNF_49_HEADERS.index("pmark")
        driver_idx = main.GR_BNF_49_HEADERS.index("Driver Name *")
        mobile_idx = main.GR_BNF_49_HEADERS.index("Mobile no   *")
        igst_idx = main.GR_BNF_49_HEADERS.index("igst")
        grdt_idx = main.GR_BNF_49_HEADERS.index("grdt *")

        self.assertEqual(str(row[grno_idx]), "33306")
        self.assertEqual(str(row[pmark_idx]), "002765")
        self.assertIn(row[driver_idx], (None, "", "None"))
        self.assertIn(row[mobile_idx], (None, "", "None"))
        self.assertEqual(str(row[igst_idx]), "0.00")
        self.assertEqual(str(row[grdt_idx]), "29/09/2026")

    def test_19_workbook_generation_with_three_brands(self):
        """
        Verify multi-brand Excel generation with Himalaya, Patanjali, and Unicharm (Vuicharm).
        Ensures all 3 formats sit side-by-side cleanly in the standard BNF 49-column template.
        """
        doc_himalaya = main.ParsedDocument(
            document_title="Himalaya_LR_32978",
            data={
                "LR No *": "32978",
                "grno (*)": "32978",
                "grdt *": "30/09/2026",
                "Customer Name *": "Himalaya Wellness Company",
                "Vehicle No *": "GJ05BZ0237",
                "pmark": "00153",
            },
            key_value_pairs=[],
            confidence_notes=""
        )
        doc_patanjali = main.ParsedDocument(
            document_title="Patanjali_LR_21349",
            data={
                "LR No *": "21349",
                "grno (*)": "21349",
                "grdt *": "05/10/2026",
                "Customer Name *": "PATANJALI FOODS LTD BISCUIT (Formerly RSIL)",
                "Vehicle No *": "RJ32GC5386",
                "pmark": "",
                "freight  *": "0",
            },
            key_value_pairs=[],
            confidence_notes=""
        )
        doc_unicharm = main.ParsedDocument(
            document_title="Unicharm_LR_32298",
            data={
                "LR No *": "32298",
                "grno (*)": "32298",
                "grdt *": "27/09/2026",
                "Customer Name *": "UNICHARM INDIA PVT. LTD.",
                "Vehicle No *": "HR55AX5490",
                "pmark": "03386",
                "freight  *": "0",
            },
            key_value_pairs=[],
            confidence_notes=""
        )

        _, user_bytes, _, _ = main._build_workbook_for_docs(
            [doc_himalaya, doc_patanjali, doc_unicharm],
            format_type="gr_bnf_49",
            user_name="Multi-Brand Enterprise QA",
            include_audit_meta=False
        )

        wb = openpyxl.load_workbook(io.BytesIO(user_bytes))
        ws = wb.active

        header_vals = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        self.assertEqual(len(header_vals), 49)

        grno_col = header_vals.index("grno (*)") + 1
        pmark_col = header_vals.index("pmark") + 1
        cust_col = header_vals.index("Customer Name *") + 1

        # Row 2 (Himalaya)
        self.assertEqual(str(ws.cell(row=2, column=grno_col).value), "32978")
        self.assertEqual(str(ws.cell(row=2, column=pmark_col).value), "00153")
        self.assertEqual(str(ws.cell(row=2, column=cust_col).value), "Himalaya Wellness Company")

        # Row 3 (Patanjali)
        self.assertEqual(str(ws.cell(row=3, column=grno_col).value), "21349")
        self.assertIn(ws.cell(row=3, column=pmark_col).value, (None, "", "None"))
        self.assertEqual(str(ws.cell(row=3, column=cust_col).value), "PATANJALI FOODS LTD BISCUIT (Formerly RSIL)")

        # Row 4 (Unicharm)
        self.assertEqual(str(ws.cell(row=4, column=grno_col).value), "32298")
        self.assertEqual(str(ws.cell(row=4, column=pmark_col).value), "03386")
        self.assertEqual(str(ws.cell(row=4, column=cust_col).value), "UNICHARM INDIA PVT. LTD.")


if __name__ == "__main__":
    unittest.main(verbosity=2)
